import docx
from docx.oxml.ns import qn
from typing import List

from models.document_block import BlockType, DocumentBlock

from modules.document import BlockParsedType, BlockClassifiedType

class DocxParser:
    def parse(self, file_path: str) -> List[DocumentBlock]:
        doc = docx.Document(file_path)
        blocks: List[DocumentBlock] = []
        table_index = 0
        toc_index = 0

        for element in doc.element.body:
            tag_name = element.tag.split("}")[-1]

            if tag_name == "p":
                if element.xpath('./ancestor::*[local-name()="sdt"]'):
                    continue

                para = next((p for p in doc.paragraphs if p._element == element), None)
                if not para:
                    continue

                text = para.text.strip()
                if not text:
                    continue

                block_type = BlockType.TEXT
                if para._element.xpath("./w:pPr/w:numPr"):
                    block_type = BlockType.LIST_ITEM
                    val_attr = para._element.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
                    list_level = int(val_attr[0]) if val_attr else 0
                    marker = ["•", "*", "+", "-"][list_level] if list_level < 4 else "•"
                    text = f"{marker} {text}"

                blocks.append(DocumentBlock(content=text, type=block_type, id=f"par_{para._element.get(qn('w14:paraId'))}"))

            elif tag_name == "tbl":
                if element.xpath('./ancestor::*[local-name()="sdt"]'):
                    continue

                table = next((t for t in doc.tables if t._element == element), None)
                if table:
                    table_rows = []
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip():
                            table_rows.append(row_text)

                    if table_rows:
                        blocks.append(DocumentBlock(content="\n".join(table_rows), type=BlockType.TABLE, id=f"tbl_{table_index}"))
                        table_index += 1

            elif tag_name == "sdt":
                is_toc = bool(element.xpath('.//*[local-name()="hyperlink"] or .//*[local-name()="fldSimple"][contains(@*[local-name()="instr"], "TOC")]'))
                if is_toc:
                    toc_rows = []
                    for p in element.xpath('.//*[local-name()="p"]'):
                        text_nodes = p.xpath('.//*[local-name()="r"]/*[local-name()="t"]')
                        p_text = "".join(node.text for node in text_nodes if node.text).strip()
                        if p_text:
                            toc_rows.append(p_text)

                    if toc_rows:
                        blocks.append(DocumentBlock(content="\n".join(toc_rows), type=BlockType.TABLE, id=f"toc_{toc_index}"))
                        toc_index += 1

        return blocks
        